"""Ephemeral attachment parsing for the operator chat.

Files are decoded and summarized only for the current request.  The API does
not write their bytes to disk or to the ResolveOps database.
"""
from __future__ import annotations

import base64
import binascii
import re
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Any


MAX_FILES = 5
MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_TOTAL_BYTES = 12 * 1024 * 1024
MAX_TEXT_CHARS_PER_FILE = 24_000
MAX_TEXT_CHARS_TOTAL = 60_000

TEXT_EXTENSIONS = {
    '.txt', '.md', '.csv', '.tsv', '.json', '.log', '.xml', '.html', '.htm',
    '.yaml', '.yml', '.ini', '.cfg', '.py', '.js', '.ts', '.sql', '.rst',
}
IMAGE_MEDIA_TYPES = {
    '.png': 'image/png', '.jpg': 'image/jpeg', '.jpeg': 'image/jpeg', '.webp': 'image/webp',
}


class AttachmentError(ValueError):
    """Raised for an attachment the chat deliberately does not accept."""


@dataclass(frozen=True)
class ChatAttachment:
    filename: str
    content_type: str
    size_bytes: int
    text: str = ''
    image_data_url: str | None = None

    def audit_metadata(self) -> dict[str, Any]:
        return {
            'filename': self.filename,
            'content_type': self.content_type,
            'size_bytes': self.size_bytes,
            'kind': 'image' if self.image_data_url else 'document',
        }


def decode_attachments(items: list[dict[str, str]]) -> list[ChatAttachment]:
    if not items:
        return []
    if len(items) > MAX_FILES:
        raise AttachmentError(f'一次最多上传 {MAX_FILES} 个文件。')
    attachments: list[ChatAttachment] = []
    total_bytes = 0
    total_text = 0
    for item in items:
        filename = Path(str(item.get('filename') or '')).name.strip()
        if not filename:
            raise AttachmentError('附件缺少文件名。')
        try:
            raw = base64.b64decode(str(item.get('data_base64') or ''), validate=True)
        except (ValueError, binascii.Error) as exc:
            raise AttachmentError(f'附件 {filename} 无法读取。') from exc
        if not raw:
            raise AttachmentError(f'附件 {filename} 为空。')
        if len(raw) > MAX_FILE_BYTES:
            raise AttachmentError(f'附件 {filename} 超过 5 MB 限制。')
        total_bytes += len(raw)
        if total_bytes > MAX_TOTAL_BYTES:
            raise AttachmentError('本次附件总大小超过 12 MB 限制。')
        attachment = _parse_attachment(filename, str(item.get('content_type') or ''), raw)
        if attachment.text:
            remaining = MAX_TEXT_CHARS_TOTAL - total_text
            if remaining <= 0:
                attachment = ChatAttachment(
                    filename=attachment.filename, content_type=attachment.content_type,
                    size_bytes=attachment.size_bytes, text='[内容因本轮附件总长度限制而未发送]',
                )
            elif len(attachment.text) > remaining:
                attachment = ChatAttachment(
                    filename=attachment.filename, content_type=attachment.content_type,
                    size_bytes=attachment.size_bytes,
                    text=f'{attachment.text[:remaining]}\n[内容已截断]',
                )
            total_text += len(attachment.text)
        attachments.append(attachment)
    return attachments


def parse_attachment_bytes(filename: str, content_type: str, raw: bytes) -> ChatAttachment:
    """Parse trusted local bytes after the local-file tool has authorized a path."""
    if len(raw) > MAX_FILE_BYTES:
        raise AttachmentError(f'附件 {filename} 超过 5 MB 限制。')
    return _parse_attachment(Path(filename).name, content_type, raw)


def _parse_attachment(filename: str, content_type: str, raw: bytes) -> ChatAttachment:
    suffix = Path(filename).suffix.lower()
    if suffix in IMAGE_MEDIA_TYPES:
        media_type = IMAGE_MEDIA_TYPES[suffix]
        encoded = base64.b64encode(raw).decode('ascii')
        return ChatAttachment(filename, media_type, len(raw), image_data_url=f'data:{media_type};base64,{encoded}')
    if suffix in TEXT_EXTENSIONS:
        return ChatAttachment(filename, content_type or 'text/plain', len(raw), text=_clean_text(_decode_text(raw)))
    if suffix == '.pdf':
        return ChatAttachment(filename, 'application/pdf', len(raw), text=_clean_text(_read_pdf(raw, filename)))
    if suffix == '.docx':
        return ChatAttachment(filename, 'application/vnd.openxmlformats-officedocument.wordprocessingml.document', len(raw), text=_clean_text(_read_docx(raw, filename)))
    if suffix == '.xlsx':
        return ChatAttachment(filename, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', len(raw), text=_clean_text(_read_xlsx(raw, filename)))
    raise AttachmentError('不支持该文件类型。可上传文本、CSV、JSON、PDF、Word、Excel 或 PNG/JPEG/WebP 图片。')


def _decode_text(raw: bytes) -> str:
    for encoding in ('utf-8-sig', 'gb18030'):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise AttachmentError('文本文件不是 UTF-8 或 GB18030 编码。')


def _read_pdf(raw: bytes, filename: str) -> str:
    try:
        from pypdf import PdfReader
        reader = PdfReader(BytesIO(raw))
        parts: list[str] = []
        for page in reader.pages:
            parts.append(page.extract_text() or '')
            if sum(len(part) for part in parts) >= MAX_TEXT_CHARS_PER_FILE:
                break
        return '\n'.join(parts) or '[未从 PDF 中提取到可读文字；扫描件请上传页面图片。]'
    except ImportError as exc:
        raise AttachmentError(f'当前运行时未安装 PDF 解析组件，无法读取 {filename}。') from exc
    except Exception as exc:
        raise AttachmentError(f'无法读取 PDF：{filename}。') from exc


def _read_docx(raw: bytes, filename: str) -> str:
    try:
        from docx import Document
        document = Document(BytesIO(raw))
        parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            parts.extend(' | '.join(cell.text.strip() for cell in row.cells) for row in table.rows)
        return '\n'.join(parts) or '[Word 文档中没有可读文字。]'
    except ImportError as exc:
        raise AttachmentError(f'当前运行时未安装 Word 解析组件，无法读取 {filename}。') from exc
    except Exception as exc:
        raise AttachmentError(f'无法读取 Word 文档：{filename}。') from exc


def _read_xlsx(raw: bytes, filename: str) -> str:
    try:
        from openpyxl import load_workbook
        workbook = load_workbook(BytesIO(raw), read_only=True, data_only=True)
        parts: list[str] = []
        for sheet in workbook.worksheets[:5]:
            parts.append(f'[工作表：{sheet.title}]')
            for row_index, row in enumerate(sheet.iter_rows(values_only=True)):
                if row_index >= 500:
                    parts.append('[该工作表其余行已截断]')
                    break
                values = [str(value).strip() if value is not None else '' for value in row]
                if any(values):
                    parts.append('\t'.join(values))
                if sum(len(part) for part in parts) >= MAX_TEXT_CHARS_PER_FILE:
                    break
        return '\n'.join(parts) or '[Excel 文件中没有可读单元格。]'
    except ImportError as exc:
        raise AttachmentError(f'当前运行时未安装 Excel 解析组件，无法读取 {filename}。') from exc
    except Exception as exc:
        raise AttachmentError(f'无法读取 Excel 文件：{filename}。') from exc


def _clean_text(value: str) -> str:
    value = value.replace('\x00', '').strip()
    value = re.sub(r'(?im)\b(api[_ -]?key|secret|password|passwd|token)\s*([:=])\s*([^\s,;]+)', r'\1\2[已隐藏]', value)
    value = re.sub(r'(?i)\bBearer\s+[A-Za-z0-9._~+\-/=]{12,}', 'Bearer [已隐藏]', value)
    value = re.sub(r'(?i)(://[^\s:/]+:)[^@\s/]+(@)', r'\1[已隐藏]\2', value)
    if len(value) > MAX_TEXT_CHARS_PER_FILE:
        value = f'{value[:MAX_TEXT_CHARS_PER_FILE]}\n[内容已截断]'
    return value or '[文件没有可读文本。]'
